"""
SpreadsheetReader - Biblioteca para leitura de arquivos XLSX
Lê planilhas de ciclos de testes e retorna dados estruturados para carga de pedidos
"""

import openpyxl
from pathlib import Path
import re


class SpreadsheetReader:
    """Classe para ler e processar dados de planilhas XLSX de ciclos de testes"""

    ROBOT_LIBRARY_SCOPE = 'GLOBAL'

    def __init__(self):
        self.workbook = None
        self.worksheet = None

    def read_ciclo_de_testes(self, ciclo_id):
        """
        Lê arquivo de ciclo de testes e retorna lista de casos de teste.
        
        Args:
            ciclo_id: ID do ciclo (ex: 369 para ciclo_de_testes_369.xlsx)
            
        Returns:
            Lista de dicionários contendo dados dos pedidos:
            [
                {
                    'id_cenario': str,
                    'cpf_cnpj': str ou None,
                    'produtos': [
                        {'id': str, 'tipo_entrega': str},
                        {'id': str, 'tipo_entrega': str}
                    ],
                    'forma_pagamento': str
                },
                ...
            ]
        """
        # Construir caminho do arquivo
        file_path = Path(__file__).parent.parent / "data" / f"ciclo_de_testes_{ciclo_id}.xlsx"
        
        if not file_path.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")
        
        # Carregar workbook
        self.workbook = openpyxl.load_workbook(str(file_path))
        self.worksheet = self.workbook.active
        
        casos_teste = []
        
        # Iterar sobre linhas (começando da linha 2, pois linha 1 é cabeçalho)
        for row_idx in range(2, self.worksheet.max_row + 1):
            # Obter valores de cada coluna
            id_cenario = self._normalize_value(self._get_cell_value(self.worksheet.cell(row_idx, 1)))
            
            # CPF/CNPJ: converter float para int antes de normalizar
            cpf_cnpj_raw = self._get_cell_value(self.worksheet.cell(row_idx, 2))
            if isinstance(cpf_cnpj_raw, float):
                cpf_cnpj = self._normalize_value(str(int(cpf_cnpj_raw)))
            else:
                cpf_cnpj = self._normalize_value(cpf_cnpj_raw)
            
            id_ret_1 = self._get_cell_value(self.worksheet.cell(row_idx, 3))
            tipo_ent_1 = self._normalize_value(self._get_cell_value(self.worksheet.cell(row_idx, 4)))
            id_ret_2 = self._get_cell_value(self.worksheet.cell(row_idx, 5))
            tipo_ent_2 = self._normalize_value(self._get_cell_value(self.worksheet.cell(row_idx, 6)))
            forma_pagamento = self._normalize_value(self._get_cell_value(self.worksheet.cell(row_idx, 7)))
            
            # Pular linhas vazias
            if not id_cenario:
                continue
            
            # Construir lista de produtos
            produtos = []
            if id_ret_1:
                produtos.append({
                    'id': str(int(id_ret_1)) if isinstance(id_ret_1, float) else str(id_ret_1),
                    'tipo_entrega': tipo_ent_1
                })
            
            if id_ret_2:
                produtos.append({
                    'id': str(int(id_ret_2)) if isinstance(id_ret_2, float) else str(id_ret_2),
                    'tipo_entrega': tipo_ent_2
                })
            
            # Construir caso de teste
            caso = {
                'id_cenario': id_cenario,
                'cpf_cnpj': cpf_cnpj if cpf_cnpj and cpf_cnpj.lower() != 'null' else None,
                'produtos': produtos,
                'forma_pagamento': forma_pagamento
            }
            
            casos_teste.append(caso)
        
        self.workbook.close()
        return casos_teste

    def _get_cell_value(self, cell):
        """
        Extrai valor da célula, tratando 'null' e valores vazios.
        
        Args:
            cell: Célula openpyxl
            
        Returns:
            Valor da célula ou None se vazia/null
        """
        value = cell.value
        
        # Tratar None e strings vazias
        if value is None or value == '':
            return None
        
        # Tratar string 'null'
        if isinstance(value, str) and value.strip().lower() == 'null':
            return None
        
        return value

    def _normalize_value(self, value):
        """
        Normaliza valor: remove espaços, caracteres invisíveis, quebras de linha.
        
        Args:
            value: Valor a normalizar (string ou outro tipo)
            
        Returns:
            Valor normalizado ou None se vazio
        """
        if value is None:
            return None
        
        # Converter para string se não for
        if not isinstance(value, str):
            value = str(value)
        
        # Remove espaços em branco no início e fim
        value = value.strip()
        
        # Remove caracteres de controle (invisíveis)
        # Mantém apenas caracteres visíveis: letras, números, espaço, hífens, etc
        value = ''.join(char for char in value if ord(char) >= 32 or char in '\n\r\t')
        
        # Remove quebras de linha e substitui por espaço
        value = re.sub(r'[\n\r\t]+', ' ', value)
        
        # Remove espaços múltiplos
        value = re.sub(r'\s+', ' ', value).strip()
        
        # Retorna None se ficou vazio
        return value if value else None


